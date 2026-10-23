#!/usr/bin/env python3
"""把原版十五年厂长事件写入管理员 Excel，保留原工作簿样式。"""
from __future__ import annotations

import json
import os
import sys
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, PatternFill


ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "游戏数据配置.xlsx"


def clone_cell_style(source, target):
    target._style = copy(source._style)
    if source.has_style:
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy(source.protection)


def append_parameter(ws, parameter_id: str, name: str, value: float, unit: str, note: str):
    id_column = 2
    for row in range(1, ws.max_row + 1):
        if str(ws.cell(row, id_column).value or "").strip() == parameter_id:
            ws.cell(row, 4).value = value
            return
    template_row = ws.max_row
    while template_row > 1 and not ws.cell(template_row, id_column).value:
        template_row -= 1
    target_row = template_row + 1
    for column in range(1, ws.max_column + 1):
        clone_cell_style(ws.cell(template_row, column), ws.cell(target_row, column))
    values = ["事件规则", parameter_id, name, value, unit, "不换算", "管理员配置", "立即生效", note]
    for column, cell_value in enumerate(values, 1):
        ws.cell(target_row, column).value = cell_value


def main():
    if len(sys.argv) != 2:
        raise SystemExit("用法：python import_legacy_events.py <事件JSON>")
    events = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    workbook = load_workbook(BOOK)
    ws = workbook["季度事件"]
    headers = [
        "事件ID", "标题", "触发条件", "出现权重", "是否一次性", "日志正文",
        "资金变化_万元", "压力变化", "好感对象ID", "好感变化", "关联项目ID", "启用", "备注",
        "事件类型", "年份", "季节", "说话人ID", "说话人", "选项JSON", "未触发文本", "原版条件",
    ]
    ws.cell(3, 1).value = (
        "本表是季度特殊事件数据库。年度事件在对应年份和季节的季度末参与抽取；每季度最多出现一条，"
        "是否出现由经营参数 yearEventChance 控制。可直接修改标题、正文、权重、选项和启用状态。"
    )
    for column, header in enumerate(headers, 1):
        if column > ws.max_column:
            clone_cell_style(ws.cell(5, min(13, ws.max_column)), ws.cell(5, column))
            ws.column_dimensions[ws.cell(5, column).column_letter].width = 18
        ws.cell(5, column).value = header

    template_row = 6
    template_cells = [copy(ws.cell(template_row, col)._style) for col in range(1, max(ws.max_column, len(headers)) + 1)]
    if ws.max_row >= template_row:
        ws.delete_rows(template_row, ws.max_row - template_row + 1)

    editable_fill = PatternFill("solid", fgColor="FFF1C9")
    for index, event in enumerate(events, template_row):
        values = [
            event["id"], event["title"], f"第{event['year']}年·{event['season']}季", 100, "是", event["text"],
            0, 0, "", 0, "", "否" if event.get("build") else "是",
            ("原版年度技改菜单，仅作为资料保留；按已确认规则不在当前游戏触发。" if event.get("build") else
             "源自 green-factory-15-years 厂长路线；原版条件事件已保留条件文本，当前按年份与季节抽取。"),
            "year_story", event["year"], event["season"], event["speakerId"], event["speaker"],
            json.dumps(event["choices"], ensure_ascii=False), event["missText"], event["legacyCondition"],
        ]
        for column, value in enumerate(values, 1):
            cell = ws.cell(index, column)
            if column - 1 < len(template_cells):
                cell._style = copy(template_cells[column - 1])
            cell.value = value
            cell.alignment = Alignment(
                horizontal=cell.alignment.horizontal,
                vertical="top",
                text_rotation=cell.alignment.text_rotation,
                wrap_text=True,
                shrink_to_fit=cell.alignment.shrink_to_fit,
                indent=cell.alignment.indent,
            )
            if column in {2, 4, 6, 12, 15, 16, 18, 19, 20}:
                cell.fill = copy(editable_fill)
        ws.row_dimensions[index].height = 42
    ws.auto_filter.ref = f"A5:U{ws.max_row}"
    ws.freeze_panes = "A6"
    widths = {1: 29, 2: 25, 3: 18, 4: 11, 5: 11, 6: 52, 12: 10, 13: 38, 14: 14, 15: 9, 16: 9, 17: 13, 18: 13, 19: 70, 20: 42, 21: 52}
    for column, width in widths.items():
        ws.column_dimensions[ws.cell(5, column).column_letter].width = width

    append_parameter(
        workbook["经营参数"], "yearEventChance", "季度末年度事件出现概率", 0.5, "0-1",
        "每季度末先判定是否出现年度事件；0 表示不出现，1 表示必定出现。",
    )
    if workbook.calculation is not None:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
    temporary = BOOK.with_suffix(".xlsx.tmp")
    workbook.save(temporary)
    os.replace(temporary, BOOK)
    print(f"已写入 {len(events)} 条年度季度事件：{BOOK}")


if __name__ == "__main__":
    main()
